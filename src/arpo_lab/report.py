"""Measured comparison and a Word report skeleton; never invent algorithm results."""

import csv
import html
import io
import json
import statistics
import tomllib
from datetime import datetime
from pathlib import Path

from .storage import read_json, save_json


def load_runs(paths):
    runs = []
    for value in paths:
        path = Path(value).resolve()
        if path.is_file():
            path = path.parent
        if (path / "experiment.json").exists():
            manifest = read_json(path / "experiment.json")
            for entry in manifest["runs"]:
                run = read_json(path / entry["path"] / "run.json")
                run["source_path"] = str(path / entry["path"])
                run["repeat"] = entry["repeat"]
                runs.append(run)
        else:
            run = read_json(path / "run.json")
            run["source_path"] = str(path)
            runs.append(run)
    return runs


def elapsed_with_wait(run):
    if not run.get("finished_at"):
        return None
    return (
        datetime.fromisoformat(run["finished_at"]) - datetime.fromisoformat(run["created_at"])
    ).total_seconds()


def summaries(runs):
    groups = {}
    for run in runs:
        groups.setdefault(run["label"], []).append(run)
    result = {}
    for label, group in groups.items():
        exact = all(r["metrics"]["tokens_complete"] for r in group)
        successful = [r for r in group if r["status"] == "passed"]
        rows = {
            "runs": len(group),
            "passed": len(successful),
            "tokens_complete": exact,
            "initial_prompt_chars_median": statistics.median([r["initial_prompt_chars"] for r in group]),
            "user_followups_median": statistics.median([r["user_followups"] for r in group]),
            "statuses": [r["status"] for r in group],
            "successful_active_time_median_s": statistics.median(
                [r["metrics"]["active_time_s"] for r in successful]
            )
            if successful
            else None,
        }
        for key in [
            "active_time_s",
            "model_time_s",
            "validation_time_s",
            "requests",
            "tool_calls",
            "total_tokens_known",
            "prompt_tokens_known",
            "completion_tokens_known",
            "repair_cycles",
        ]:
            values = [r["metrics"][key] for r in group]
            rows[key + "_median"] = statistics.median(values)
            rows[key + "_min"] = min(values)
            rows[key + "_max"] = max(values)
        elapsed = [value for r in group if (value := elapsed_with_wait(r)) is not None]
        rows["elapsed_with_wait_s_median"] = statistics.median(elapsed) if elapsed else None
        rows["elapsed_with_wait_s_min"] = min(elapsed) if elapsed else None
        rows["elapsed_with_wait_s_max"] = max(elapsed) if elapsed else None
        result[label] = rows
    return result


def comparison_rows(summary):
    labels = list(summary)
    rows = [["Показатель", *labels]]
    rows.append(
        ["Успешные / всего", *[f"{summary[label]['passed']} / {summary[label]['runs']}" for label in labels]]
    )
    for caption, key in [
        ("Размер промпта, символы", "initial_prompt_chars"),
        ("Уточнения пользователя", "user_followups"),
        ("Активное время, с", "active_time_s"),
        ("Календарное время с паузами, с", "elapsed_with_wait_s"),
        ("Время API, с", "model_time_s"),
        ("Время проверок, с", "validation_time_s"),
        ("Запросы HTTP", "requests"),
        ("Входные токены", "prompt_tokens_known"),
        ("Выходные токены", "completion_tokens_known"),
        ("Всего токенов", "total_tokens_known"),
        ("Циклы исправления", "repair_cycles"),
    ]:
        row = [caption]
        for label in labels:
            group = summary[label]
            median = group[key + "_median"]
            text = "—" if median is None else f"{median:.2f}" if key.endswith("_s") else f"{median:g}"
            if "tokens" in key and not group["tokens_complete"]:
                text += "*"
            row.append(text)
        rows.append(row)
    return rows


def md_table(rows):
    def cell(value):
        return str(value).replace("|", "\\|").replace("\n", " ")

    lines = ["| " + " | ".join(cell(c) for c in row) + " |" for row in rows]
    lines.insert(1, "| " + " | ".join("---" for _ in rows[0]) + " |")
    return "\n".join(lines)


def comparison_warnings(runs):
    warnings = []
    if any(not r["metrics"]["tokens_complete"] for r in runs):
        warnings.append(
            "Есть запросы без usage: токены и стоимость неполны и показаны как известная часть (*)."
        )
    if any(not r["metrics"].get("timing_complete", True) for r in runs):
        warnings.append("Время одного из запусков неполно после аварийного завершения процесса.")
    if any(not r["metrics"].get("requests_complete", True) for r in runs):
        warnings.append("Число запросов может быть неполным после аварийного завершения процесса.")
    if any(r["status"] != "passed" for r in runs):
        warnings.append(
            "Не все реализации прошли проверки. "
            "Нельзя сравнивать их стоимость как стоимость успешного решения."
        )
    settings = [
        json.dumps(
            {k: r["config"][k] for k in ["model", "agent", "tools", "validation", "pricing"]},
            sort_keys=True,
            ensure_ascii=False,
        )
        for r in runs
    ]
    if len(set(settings)) > 1:
        warnings.append("Параметры запусков отличаются: различие результата нельзя приписать только промпту.")
    if len({r["seed_sha256"] for r in runs}) > 1:
        warnings.append("Начальные файлы отличаются.")
    if len({json.dumps(r.get("model_revision"), sort_keys=True) for r in runs}) > 1:
        warnings.append("Ревизии весов модели отличаются.")
    commands = [
        r.get("server_metadata", {}).get("command") if r.get("server_metadata") else None for r in runs
    ]
    if len({json.dumps(command) for command in commands}) > 1:
        warnings.append("Команды запуска сервера отличаются или известны не для всех запусков.")
    for label in {r["label"] for r in runs}:
        if len({r["initial_prompt_sha256"] for r in runs if r["label"] == label}) > 1:
            warnings.append(f"Промпт варианта {label} менялся между повторами.")
    return warnings


NOTE_SECTIONS = [
    ("task", "Постановка задачи"),
    ("algorithm", "Алгоритм"),
    ("complexity", "Оценка сложности"),
    ("improvements", "Улучшения"),
    ("testing", "Результаты тестирования алгоритма"),
    ("integration", "Поддержка компонентов ACOM"),
    ("visualization", "Визуализация"),
    ("conclusion", "Выводы"),
]


def write_report(runs, output, formats=("md", "html", "csv", "json"), notes_path=None, benchmarks=None):
    if not runs:
        raise ValueError("No runs to report")
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    notes = {}
    if notes_path:
        with Path(notes_path).open("rb") as stream:
            notes = tomllib.load(stream)
    benchmark_rows = []
    if benchmarks:
        with Path(benchmarks).open(encoding="utf-8-sig", newline="") as stream:
            benchmark_rows = list(csv.reader(stream))
        if benchmark_rows and (
            len(benchmark_rows[0]) > 7 or any(len(r) != len(benchmark_rows[0]) for r in benchmark_rows)
        ):
            raise ValueError("Benchmark CSV must have consistent rows and at most 7 columns")
    summary = summaries(runs)
    rows = comparison_rows(summary)
    warnings = comparison_warnings(runs)
    meta = {**runs[0]["config"]["report"], **notes.get("report", {})}
    title = meta["title"]
    lines = [
        f"# {title}",
        "",
        f"Автор: {meta.get('author') or '[заполнить]'}. Группа: {meta.get('group') or '[заполнить]'}.",
        "",
        (
            "Сравниваются решения одной задачи при кратком и подробном промпте. "
            if len(summary) > 1
            else "Представлены результаты варианта " + next(iter(summary)) + ". "
        )
        + "Таблица содержит медианы всех запусков. "
        + "Работоспособность означает только прохождение заданных проверок.",
        "",
    ]
    for warning in warnings:
        lines.extend(["Внимание: " + warning, ""])
    for key, heading in NOTE_SECTIONS[:3]:
        lines.extend(["## " + heading, "", notes.get(key, "[Заполнить по проверенной реализации.]"), ""])
    lines.extend(
        [
            "## Сравнение работы ИИ",
            "",
            md_table(rows),
            "",
            "Входные токены суммируются по всем запросам вместе с повторно переданной историей "
            "и системным промптом. "
            "Ошибки сети учитываются как попытки HTTP; их токены могут быть неизвестны. "
            "Активное время включает API, инструменты, проверки и повторы, "
            "но исключает паузы между сессиями. "
            "Время выполнения лабораторного алгоритма измеряется отдельно.",
            "",
        ]
    )
    raw_rows = [["ID", "Вариант", "Статус", "Активное время с", "HTTP", "Токены", "Уточнения пользователя"]]
    for run in runs:
        m = run["metrics"]
        raw_rows.append(
            [
                run["id"],
                run["label"],
                run["status"],
                f"{m['active_time_s']:.2f}",
                m["requests"],
                str(m["total_tokens_known"]) + ("" if m["tokens_complete"] else "*"),
                run["user_followups"],
            ]
        )
    lines.extend(["### Все запуски", "", md_table(raw_rows), ""])
    for key, heading in NOTE_SECTIONS[3:]:
        lines.extend(["## " + heading, "", notes.get(key, "[Заполнить по проверенной реализации.]"), ""])
        if key == "testing" and benchmark_rows:
            lines.extend(
                [
                    "Данные из предоставленного CSV. Единицы и методику задаёт автор измерений.",
                    "",
                    md_table(benchmark_rows),
                    "",
                ]
            )
    lines.extend(["## Воспроизводимость", ""])
    seen_settings = set()
    for run in runs:
        model = run["config"]["model"]
        settings = (
            f"Модель {model['name']}; temperature={model['temperature']}; "
            f"top_p={model['top_p']}; seed={model['seed']}; max_tokens={model['max_tokens']}. "
            f"Протокол {run['config']['agent']['protocol']}."
        )
        if model.get("extra_body"):
            settings += " Дополнительные параметры API: " + json.dumps(
                model["extra_body"], ensure_ascii=False, sort_keys=True
            ) + "."
        if run.get("seed_context_chars"):
            settings += (
                f" Общий контекст шаблона: {run['seed_context_chars']} символов (входит во входные токены)."
            )
        if run.get("model_revision"):
            settings += f" Ревизия весов: `{run['model_revision']['revision']}`."
        if run.get("sdk_metadata"):
            sdk = run["sdk_metadata"]
            settings += f" Eco SDK: {sdk['archive']}; SHA256 архива: `{sdk['archive_sha256']}`."
        if run.get("server_metadata"):
            settings += " Сохранённая команда сервера: " + " ".join(run["server_metadata"]["command"]) + "."
        if settings not in seen_settings:
            lines.extend([settings, ""])
            seen_settings.add(settings)
    parents = {str(Path(run["source_path"]).parent) for run in runs}
    if len(parents) == 1 and all(Path(run["source_path"]).name == run["id"] for run in runs):
        lines.extend(
            [
                "Каталог результатов: " + next(iter(parents)) + ". "
                "Подкаталоги совпадают с ID в таблице всех запусков. "
                "Исходные промпты находятся в prompt.txt, уточнения в followup-*.txt.",
                "",
            ]
        )
        seen_prompts = set()
        for run in runs:
            prompt_key = (run["label"], run["initial_prompt_sha256"])
            if prompt_key not in seen_prompts:
                lines.append(f"- {prompt_key[0]}: промпт SHA256 `{prompt_key[1]}`.")
                seen_prompts.add(prompt_key)
    else:
        for run in runs:
            lines.append(
                f"- `{run['id']}`: промпт SHA256 `{run['initial_prompt_sha256']}`; "
                f"файлы `{run['source_path']}`."
            )
    lines.extend(
        [
            "",
            "Полная конфигурация, версии окружения, запросы, ответы, проверки "
            "и исходники сохранены рядом с run.json. "
            "Стоимость по тарифам не включает электричество и амортизацию ноутбука. "
            "При нулевых тарифах локального "
            "сервера плата провайдеру равна нулю; объём вычислений отражают токены и время.",
            "",
        ]
    )
    if notes.get("references"):
        lines.extend(["## Источники", "", *["- " + s for s in notes["references"]], ""])
    lines.extend(["## Промпты и пользовательские уточнения", ""])
    seen_prompts = set()
    for run in runs:
        key = (run["label"], run["initial_prompt_sha256"])
        if key not in seen_prompts:
            original = run["messages"][1]["content"]
            lines.extend([f"### Исходный промпт: {run['label']}", "", *original.splitlines(), ""])
            seen_prompts.add(key)
        for index in range(1, run["user_followups"] + 1):
            followup = Path(run["source_path"]) / f"followup-{index:02d}.txt"
            lines.extend([f"### Уточнение {index}: {run['label']}, {run['id']}", ""])
            lines.extend(
                followup.read_text(encoding="utf-8").splitlines()
                if followup.is_file()
                else ["Файл уточнения отсутствует; проверьте историю в run.json."]
            )
            lines.append("")
    markdown = "\n".join(lines)
    outputs = []
    for fmt in formats:
        target = output / ("comparison." + fmt)
        if fmt == "md":
            target.write_text(markdown, encoding="utf-8")
        elif fmt == "json":
            save_json(
                target,
                {
                    "schema_version": 1,
                    "summary": summary,
                    "warnings": warnings,
                    "runs": runs,
                    "benchmark_rows": benchmark_rows,
                    "notes": notes,
                },
            )
        elif fmt == "csv":
            stream = io.StringIO(newline="")
            columns = [
                "id",
                "label",
                "status",
                "initial_prompt_chars",
                "initial_prompt_bytes",
                "user_followups",
                "requests",
                "successful_requests",
                "failed_requests",
                "requests_without_usage",
                "tokens_complete",
                "timing_complete",
                "requests_complete",
                "initial_request_prompt_tokens",
                "prompt_tokens_known",
                "completion_tokens_known",
                "total_tokens_known",
                "active_time_s",
                "model_time_s",
                "tool_time_s",
                "validation_time_s",
                "tool_calls",
                "repair_cycles",
                "estimated_cost_known",
                "elapsed_with_wait_s",
                "currency",
            ]
            writer = csv.DictWriter(stream, fieldnames=columns)
            writer.writeheader()
            for run in runs:
                record = {
                    **run,
                    **run["metrics"],
                    "elapsed_with_wait_s": elapsed_with_wait(run),
                    "currency": run["config"]["pricing"]["currency"],
                }
                writer.writerow({k: record.get(k) for k in columns})
            target.write_text(stream.getvalue(), encoding="utf-8-sig")
        elif fmt == "html":
            # Escape all user/model data; self-contained browser/print view.
            table = (
                "<table>"
                + "".join(
                    "<tr>"
                    + "".join(
                        f"<{'th' if i == 0 else 'td'}>{html.escape(str(c))}</{'th' if i == 0 else 'td'}>"
                        for c in row
                    )
                    + "</tr>"
                    for i, row in enumerate(rows)
                )
                + "</table>"
            )
            target.write_text(
                '<!doctype html><html lang="ru"><meta charset="utf-8">'
                + f"<title>{html.escape(title)}</title>"
                + "<style>body{max-width:1000px;margin:40px auto;padding:20px;font:16px/1.5 system-ui}"
                "table{border-collapse:collapse;width:100%}td,th{border:1px solid #ccc;padding:10px}"
                "pre{white-space:pre-wrap;overflow-wrap:anywhere}</style><body>"
                + f"<h1>{html.escape(title)}</h1>{table}<pre>{html.escape(markdown)}</pre></body></html>",
                encoding="utf-8",
            )
        elif fmt == "docx":
            write_docx(target, title, meta, lines, rows, raw_rows, benchmark_rows)
        else:
            raise ValueError(f"Unsupported format: {fmt}")
        outputs.append(target)
    return outputs


def write_docx(path, title, meta, lines, rows, raw_rows, benchmark_rows):
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("Install reports extra: pip install -e '.[reports]'") from exc
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style._element.get_or_add_rPr().find(qn("w:rFonts"))
        if fonts is not None:
            for attr in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
                fonts.attrib.pop(qn("w:" + attr), None)
    # Some python-docx templates contain an inherited title rule.
    for border in document.styles._element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1
    document.styles["Title"].font.size = Pt(20)
    document.styles["Title"].paragraph_format.space_after = Pt(12)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(8)
    document.add_paragraph(title, "Title")
    document.add_paragraph("Автор " + (meta.get("author") or "[заполнить]"))
    document.add_paragraph("Группа " + (meta.get("group") or "[заполнить]"))

    def table(values):
        if not values:
            return
        t = document.add_table(rows=0, cols=len(values[0]))
        t.autofit = False
        count = len(values[0])
        widths = [17.0] if count == 1 else [8.0] + [9.0 / (count - 1)] * (count - 1)
        if count > 3:
            widths = [17.0 / count] * count
        if values is raw_rows and count == 7:
            widths = [3.6, 1.9, 2.3, 2.4, 1.3, 2.2, 3.3]
        for column, width in zip(t.columns, widths):
            column.width = Cm(width)
        for index, row in enumerate(values):
            cells = t.add_row().cells
            for col, value in enumerate(row):
                cells[col].width = Cm(widths[col])
                cells[col].text = str(value)
                cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cells[col]._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ["top", "left", "bottom", "right"]:
                    element = OxmlElement("w:" + edge)
                    for k, v in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                        element.set(qn("w:" + k), v)
                    borders.append(element)
                tcpr.append(borders)
                if index == 0:
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), "E7EDF3")
                    tcpr.append(fill)
                for paragraph in cells[col].paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.space_before = Pt(3)
                    for run in paragraph.runs:
                        run.font.size = Pt(9 if len(row) > 3 else 11)
                        run.bold = index == 0
            if index == 0:
                header = OxmlElement("w:tblHeader")
                t.rows[-1]._tr.get_or_add_trPr().append(header)
            t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        document.add_paragraph()

    in_table = False
    table_index = 0
    for line in lines[4:]:
        if line.startswith("| "):
            if not in_table:
                table([rows, raw_rows, benchmark_rows][table_index])
                table_index += 1
            in_table = True
            continue
        in_table = False
        if not line:
            continue
        if line.startswith("## "):
            document.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("### "):
            document.add_paragraph(line[4:], "Heading 2")
        else:
            # Paths/hash appendix can contain long strings. Add break opportunities.
            text = line.replace("`", "").replace("/", "/\u200b")
            document.add_paragraph(text)
    document.save(path)
